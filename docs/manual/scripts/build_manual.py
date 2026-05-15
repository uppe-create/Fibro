from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
DOCS_DIR = ROOT / "docs" / "manual"
SCREENSHOT_DIR = DOCS_DIR / "screenshots"
OUTPUT_PATH = DOCS_DIR / "Manual_CIPF_Operacao_e_Seguranca.docx"
LOGO_PATH = ROOT / "src" / "assets" / "prefeitura-logo.png"

ACCENT = RGBColor(92, 38, 133)
INK = RGBColor(31, 24, 39)
MUTED = RGBColor(96, 87, 107)
LIGHT = RGBColor(243, 238, 248)
BORDER = RGBColor(224, 217, 233)
SUCCESS = RGBColor(26, 127, 55)
WARNING = RGBColor(180, 83, 9)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "E0D9E9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
      tc_borders = OxmlElement("w:tcBorders")
      tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        el = tc_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), color)


def apply_font(run, name="Aptos", size=11, bold=False, color: RGBColor | None = None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    rfonts = run._element.rPr.rFonts
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def style_normal(paragraph, *, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15


def add_text(paragraph, text, *, size=11, bold=False, color: RGBColor | None = None, italic=False):
    run = paragraph.add_run(text)
    apply_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def set_page(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = INK

    for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)]:
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT if style_name != "Heading 3" else INK
        style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)


def add_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.text = ""
    add_text(hp, "Manual operacional e de seguranca", size=9, color=MUTED)

    footer = section.footer
    table = footer.add_table(rows=1, cols=3, width=Inches(6.4))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    left, center, right = table.rows[0].cells
    add_text(left.paragraphs[0], "CIPF - Prefeitura de Ipero", size=8, color=MUTED)
    center.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(center.paragraphs[0], "Uso interno orientado", size=8, color=MUTED)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(right.paragraphs[0], datetime.now().strftime("%d/%m/%Y"), size=8, color=MUTED)


def add_cover(doc: Document):
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Pt(18)
    intro.add_run().add_picture(str(LOGO_PATH), width=Cm(3.8))

    kicker = doc.add_paragraph()
    style_normal(kicker, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    kicker.paragraph_format.space_before = Pt(18)
    add_text(kicker, "MANUAL OPERACIONAL", size=12, bold=True, color=ACCENT)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    add_text(title, "Carteirinha de Identificacao da Pessoa com Fibromialgia", size=25, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    style_normal(subtitle, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(subtitle, "Operacao, seguranca, LGPD e governanca do sistema CIPF.", size=12, color=MUTED)

    meta = doc.add_paragraph()
    style_normal(meta, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(meta, f"Versao do manual: {datetime.now():%d/%m/%Y}", size=10, color=MUTED)

    divider = doc.add_paragraph()
    style_normal(divider, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(divider, "______________________________________________", size=10, color=RGBColor(210, 202, 223))

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    set_cell_shading(cell, "F3EEF8")
    set_cell_border(cell, "E0D9E9")
    p = cell.paragraphs[0]
    style_normal(p, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "Documento para operadores e gestores municipais. Conteudo baseado no estado atual do sistema, com foco em uso seguro, rastreavel e auditavel.", size=11, color=INK)

    highlights = doc.add_table(rows=1, cols=3)
    highlights.alignment = WD_TABLE_ALIGNMENT.CENTER
    highlights.autofit = False
    for index, width in enumerate((Cm(5.0), Cm(5.0), Cm(5.0))):
        highlights.columns[index].width = width
    items = [
        ("Publico-alvo", "Operadores e gestores"),
        ("Escopo", "Fluxo, seguranca e LGPD"),
        ("Fonte", "Estado atual do sistema")
    ]
    for cell, (label, text) in zip(highlights.rows[0].cells, items):
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell, "E0D9E9")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = cell.paragraphs[0]
        style_normal(p1, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p1, f"{label}\n", size=10, bold=True, color=ACCENT)
        add_text(p1, text, size=10, color=INK)
    doc.add_page_break()


def add_toc(doc: Document):
    doc.add_heading("Sumario", level=1)
    toc_entries = [
        "1. Resumo executivo",
        "2. Visao geral do sistema",
        "3. Como acessar",
        "4. Operacao por modulo",
        "5. Fluxos operacionais",
        "6. Seguranca e protecao de dados",
        "7. LGPD e governanca",
        "8. Boas praticas de operacao",
        "9. Suporte e troubleshooting",
        "10. Apendices"
    ]
    for entry in toc_entries:
        p = doc.add_paragraph(style="List Bullet")
        style_normal(p, space_after=2)
        add_text(p, entry, size=10.5, color=INK)


def add_callout(doc: Document, title: str, text: str, fill="F9F6FC", title_color=ACCENT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "D7CAE5")
    p = cell.paragraphs[0]
    style_normal(p, space_after=2)
    add_text(p, f"{title}\n", size=10, bold=True, color=title_color)
    add_text(p, text, size=10.5, color=INK)


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_normal(p, space_after=3)
        add_text(p, item, size=10.5, color=INK)


def add_numbered(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        style_normal(p, space_after=3)
        add_text(p, item, size=10.5, color=INK)


def add_two_col_table(doc: Document, rows: list[tuple[str, str]], widths=(Cm(4.8), Cm(11.5))):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = widths[0]
    table.columns[1].width = widths[1]
    header = table.rows[0].cells
    for index, text in enumerate(("Item", "Descricao")):
        set_cell_shading(header[index], "F3EEF8")
        set_cell_border(header[index])
        p = header[index].paragraphs[0]
        style_normal(p, space_after=0)
        add_text(p, text, size=10, bold=True, color=ACCENT)
    for left_text, right_text in rows:
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
        p1 = cells[0].paragraphs[0]
        style_normal(p1, space_after=0)
        add_text(p1, left_text, size=10, bold=True, color=INK)
        p2 = cells[1].paragraphs[0]
        style_normal(p2, space_after=0)
        add_text(p2, right_text, size=10, color=INK)


def add_screenshot(doc: Document, filename: str, caption: str, width_cm=16.2):
    image_path = SCREENSHOT_DIR / filename
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    style_normal(cp, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cp, caption, size=9.5, italic=True, color=MUTED)


def add_module_section(doc: Document, title: str, owner: str, objective: str, actions: list[str], attention: list[str], screenshot: str, caption: str):
    doc.add_heading(title, level=2)
    add_two_col_table(doc, [
        ("Objetivo", objective),
        ("Quem usa", owner),
        ("Acoes", "; ".join(actions)),
        ("Pontos de atencao", "; ".join(attention))
    ])
    doc.add_paragraph("")
    add_screenshot(doc, screenshot, caption)


def add_section_title(doc: Document, title: str, description: str | None = None):
    doc.add_heading(title, level=1)
    if description:
        p = doc.add_paragraph()
        style_normal(p, space_after=6)
        add_text(p, description, size=10.5, color=MUTED)


def build_document() -> Document:
    doc = Document()
    set_page(doc)
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_toc(doc)

    add_section_title(doc, "Resumo executivo")
    add_bullets(doc, [
        "Sistema municipal para cadastro, analise, aprovacao, emissao, impressao e validacao publica da CIPF.",
        "Fluxo divide area publica e area interna, reduzindo exposicao de dados sensiveis.",
        "Operacao principal envolve atendente e administrador, com auditoria e restricao por papel.",
        "Seguranca central usa Supabase Auth em producao, MFA TOTP para admin, RLS e RPC publica dedicada.",
        "Documento serve para onboarding, consulta operacional e apoio a suporte e governanca."
    ])
    add_callout(doc, "Pilares de seguranca", "Separacao publico x privado, menor privilegio por papel, auditoria rastreavel, documentos em storage privado e exportacoes controladas.")

    add_section_title(doc, "Visao geral do sistema")
    add_two_col_table(doc, [
        ("Area publica", "Home institucional, validacao por QR/codigo e paginas informativas."),
        ("Area interna", "Painel, cadastros, aprovacao, impressao, auditoria, relatorios e governanca."),
        ("Perfis", "Admin: controle total; Atendente: cadastro, aprovacao e renovacao; Viewer: consulta restrita."),
        ("Status", "under_review, approved, issued, expired e cancelled."),
        ("Banco", "Supabase como base principal, com tabelas privadas, indice de CPF e validacao publica minimizada.")
    ])
    add_screenshot(doc, "home_publica.png", "Tela inicial publica do sistema CIPF.")

    add_section_title(doc, "Como acessar")
    add_numbered(doc, [
        "Abrir URL oficial do sistema.",
        "Selecionar Acessar painel para entrar na area interna.",
        "Autenticar via Supabase Auth no ambiente oficial.",
        "Se usuario admin nao concluir MFA TOTP, o acesso e reduzido para viewer.",
        "Usar apenas estacao autorizada para operacao, exportacao e impressao."
    ])
    add_callout(doc, "Ambiente de teste x producao", "Login local e aceito apenas para desenvolvimento controlado. Em producao, autenticacao segura exige Supabase Auth e MFA para administrador.", fill="FFF6E8", title_color=WARNING)
    add_screenshot(doc, "login_publico.png", "Tela publica de acesso ao painel administrativo.")

    add_section_title(doc, "Operacao por modulo", "Resumo dos modulos principais e uso esperado por perfil.")
    add_module_section(
        doc,
        "Home publica",
        "Cidadao e equipe",
        "Apresentar programa municipal, metricas agregadas e acessos principais.",
        ["Acessar painel", "Solicitar carteirinha", "Consultar FAQ"],
        ["Nao exibe dados nominais", "Metricas sao agregadas por RPC"],
        "home_publica.png",
        "Home publica com metricas minimizadas."
    )
    add_module_section(
        doc,
        "Validacao publica",
        "Cidadao e rede externa",
        "Validar autenticidade da carteirinha por QR Code ou codigo manual.",
        ["Ler QR", "Digitar codigo manual", "Visualizar dados minimos"],
        ["Nao acessa registrations diretamente", "Mostra apenas dados minimizados"],
        "validacao_publica.png",
        "Validacao publica com resposta positiva de documento ativo."
    )
    add_module_section(
        doc,
        "Configuracoes",
        "Todos os perfis",
        "Centralizar status de sessao, ambiente, autenticacao e parametros operacionais.",
        ["Ver ambiente", "Consultar permissao", "Revisar seguranca operacional"],
        ["Admin sem aal2 sofre degradacao", "Nao expor chaves nem tokens no navegador"],
        "configuracoes_admin.png",
        "Configuracoes internas com resumo de sessao e seguranca."
    )
    add_module_section(
        doc,
        "Cadastro",
        "Atendente e admin",
        "Registrar novo titular com documentos e validacoes obrigatorias.",
        ["Preencher dados", "Anexar laudo", "Enviar para under_review"],
        ["CPF unico", "CID fora de M79.7 exige justificativa", "Menor de 18 anos exige responsavel"],
        "cadastro.png",
        "Formulario principal de cadastro."
    )
    add_module_section(
        doc,
        "Pessoas",
        "Atendente e admin",
        "Consultar lista de cadastros, abrir ficha e editar dados permitidos.",
        ["Buscar por nome/CPF", "Abrir ficha", "Editar cadastro", "Acessar historico"],
        ["CPF completo depende de permissao", "Documentos sensiveis seguem trilha controlada"],
        "pessoas.png",
        "Modulo de cadastros com busca e acoes operacionais."
    )
    add_module_section(
        doc,
        "Painel",
        "Admin",
        "Monitorar fila, indicadores e cadastros recentes.",
        ["Filtrar base", "Abrir detalhes", "Exportar CSV/PDF"],
        ["Exportacao gera auditoria", "Dados sao pessoais e exigem estacao autorizada"],
        "dashboard.png",
        "Painel administrativo com indicadores e lista de registros."
    )
    add_module_section(
        doc,
        "Aprovar",
        "Atendente e admin",
        "Conferir fila e mudar cadastro para approved quando a documentacao estiver correta.",
        ["Revisar pendencias", "Aprovar cadastro", "Registrar motivo operacional"],
        ["Conferir anexos antes de aprovar", "Nao emitir neste modulo"],
        "aprovacao.png",
        "Fila operacional de aprovacao."
    )
    add_module_section(
        doc,
        "Documentos",
        "Atendente e admin",
        "Priorizar casos com pendencia documental ou inconsistencia de anexos.",
        ["Localizar pendencias", "Abrir contexto do cadastro", "Seguir para aprovacao"],
        ["Anexos sao privados", "Visualizacao sensivel deve ser auditavel"],
        "documentos.png",
        "Fila de revisao documental."
    )
    add_module_section(
        doc,
        "Retiradas",
        "Atendente e admin",
        "Controlar carteirinhas prontas para retirada e registrar entrega.",
        ["Listar prontas", "Registrar retirada", "Anotar observacao"],
        ["Retirada e evento sensivel", "Conferir identidade do titular"],
        "retiradas.png",
        "Fila de retirada de carteirinhas emitidas."
    )
    add_module_section(
        doc,
        "Relatorios",
        "Admin",
        "Gerar visoes administrativas por CSV, PDF e relatorio mensal.",
        ["Exportar CSV", "Exportar PDF", "Gerar relatorio mensal"],
        ["Nao ha backup JSON integral", "Arquivos podem conter dados pessoais"],
        "relatorios.png",
        "Area de relatorios administrativos."
    )
    add_module_section(
        doc,
        "Auditoria",
        "Admin",
        "Consultar linha do tempo de eventos, filtros e rastreabilidade.",
        ["Filtrar eventos", "Revisar severidade", "Investigar acao por cadastro"],
        ["Eventos sensiveis merecem revisao", "Auditoria vem por RPC"],
        "auditoria.png",
        "Linha do tempo de auditoria do sistema."
    )
    add_module_section(
        doc,
        "Governanca LGPD",
        "Admin",
        "Registrar pedidos do titular, incidentes, retencao e operadores ativos.",
        ["Abrir pedido LGPD", "Registrar incidente", "Revisar retencao", "Listar operadores"],
        ["Fluxo apoia governanca, nao substitui validacao juridica", "Base legal e retencao exigem alinhamento institucional"],
        "governanca.png",
        "Painel de governanca LGPD."
    )
    add_module_section(
        doc,
        "Carteirinha",
        "Admin",
        "Emitir, baixar PNG e gerar PDF tecnico para grafica.",
        ["Buscar titular", "Pre-visualizar", "Baixar PNG", "Gerar PDF para grafica"],
        ["Somente status approved ou issued", "Emissao e exportacao exigem auditoria"],
        "carteirinha.png",
        "Modulo de emissao e impressao da carteirinha."
    )

    add_section_title(doc, "Fluxos operacionais")
    flows = [
        ("Cadastrar novo caso", [
            "Abrir Novo cadastro.",
            "Preencher identificacao, contato, endereco e dados clinicos.",
            "Anexar documento oficial, comprovante, laudo e foto 3x4.",
            "Enviar cadastro para analise com status under_review."
        ]),
        ("Aprovar cadastro", [
            "Abrir modulo Aprovar.",
            "Revisar documentos e campos obrigatorios.",
            "Confirmar se laudo e comprovante estao dentro dos prazos.",
            "Mudar status para approved."
        ]),
        ("Emitir e imprimir", [
            "Abrir Carteirinha.",
            "Buscar titular por nome, CPF ou CNS.",
            "Baixar PNG da carteirinha.",
            "Usar PDF tecnico para confeccao grafica quando necessario."
        ]),
        ("Retirada e encerramento operacional", [
            "Ir para Retiradas.",
            "Selecionar cadastro emitido.",
            "Registrar retirada e observacao, quando houver.",
            "Conferir auditoria se o caso exigir rastreio."
        ])
    ]
    for title, steps in flows:
        doc.add_heading(title, level=2)
        add_numbered(doc, steps)

    add_section_title(doc, "Seguranca e protecao de dados")
    add_two_col_table(doc, [
        ("Autenticacao", "Supabase Auth em producao; login local apenas para teste controlado."),
        ("MFA", "Administrador precisa de TOTP e sessao aal2; sem isso, acesso degrada para viewer."),
        ("Autorizacao", "Permissoes centralizadas por papel e por acao sensivel."),
        ("RLS", "Banco usa Row Level Security para separar publico, viewer, atendente e admin."),
        ("Validacao publica", "RPC validate_cipf_public com dados minimos, sem leitura direta de registrations."),
        ("Documentos", "Storage privado com signed URL curta e trilha de visualizacao sensivel."),
        ("Auditoria", "Eventos operacionais e de exportacao registrados por RPC."),
        ("Exportacoes", "CSV/PDF permitidos; backup JSON integral desativado por LGPD.")
    ])
    add_callout(doc, "Separacao critica", "Area publica nao deve consultar registros privados. O desenho atual separa validacao publica, metricas agregadas e operacao interna.", fill="ECFDF3", title_color=SUCCESS)

    add_section_title(doc, "LGPD e governanca")
    add_bullets(doc, [
        "Base legal minima vinculada a politica publica, tutela da saude e controle interno.",
        "Sistema trata dados pessoais, contato, documentos e dados sensiveis de saude.",
        "Validacao publica usa minimizacao: nome, CPF mascarado, emissao, validade, status e assinatura visual.",
        "Direitos do titular e incidentes possuem modulo interno dedicado para registro e acompanhamento.",
        "Retencao ainda depende de validacao institucional formal e tabela definitiva."
    ])
    add_screenshot(doc, "governanca.png", "Exemplo do modulo de governanca LGPD com fila e cadastros de apoio.")

    add_section_title(doc, "Boas praticas de operacao")
    add_bullets(doc, [
        "Nao usar dados reais em homologacao ou treinamento.",
        "Nao compartilhar credenciais, chave TOTP ou codigos de autenticador.",
        "Conferir documentos antes de aprovar cadastro.",
        "Conferir status e validade antes de emitir ou reemitir carteirinha.",
        "Exportar somente quando houver necessidade administrativa legitima.",
        "Consultar auditoria quando houver incidente, divergencia ou acao sensivel."
    ])

    add_section_title(doc, "Suporte e troubleshooting")
    add_two_col_table(doc, [
        ("Login falhou", "Confirmar usuario, conectividade e modo de autenticacao configurado."),
        ("Admin sem acesso", "Validar MFA TOTP e sessao aal2."),
        ("Documento nao abre", "Checar permissao, storage privado e existencia do anexo."),
        ("QR nao valida", "Conferir status issued, validade e assinatura visual/codigo manual."),
        ("Perfil sem permissao", "Revisar papel no sistema e regras da matriz de permissoes."),
        ("Erro em exportacao", "Repetir em estacao autorizada e revisar logs/auditoria."),
        ("Tela estranha ou branca", "Recarregar, revisar sessao local e abrir rota segura como Configuracoes.")
    ])

    add_section_title(doc, "Apendices")
    doc.add_heading("Mapa resumido de modulos", level=2)
    add_bullets(doc, [
        "Publico: Inicio, Validar, Contato, Termos, Privacidade, Acessibilidade, Suporte.",
        "Interno admin: Painel, Cadastros, Novo cadastro, Aprovar, Documentos, Retiradas, Carteirinha, Relatorios, Auditoria, LGPD, Configuracoes.",
        "Interno atendente: Cadastros, Novo cadastro, Aprovar, Documentos, Retiradas, Configuracoes.",
        "Interno viewer: somente Configuracoes."
    ])

    doc.add_heading("Matriz resumida de permissoes", level=2)
    add_two_col_table(doc, [
        ("Administrador", "Emite, imprime, exporta, cancela, arquiva, audita e opera governanca."),
        ("Atendente", "Cadastra, edita, aprova, renova e opera filas sem impressao."),
        ("Viewer", "Consulta restrita a configuracoes.")
    ])

    doc.add_heading("Checklist de operacao segura", level=2)
    add_numbered(doc, [
        "Entrar em estacao autorizada.",
        "Confirmar perfil e MFA quando aplicavel.",
        "Conferir anexos antes de aprovar.",
        "Emitir somente cadastro apto.",
        "Registrar retirada e justificativas necessarias.",
        "Encerrar sessao ao final do uso."
    ])

    doc.add_heading("Glosario curto", level=2)
    add_two_col_table(doc, [
        ("CIPF", "Carteirinha de Identificacao da Pessoa com Fibromialgia."),
        ("MFA", "Autenticacao multifator com codigo TOTP."),
        ("RLS", "Row Level Security aplicada no banco."),
        ("RPC", "Funcao exposta pelo banco para operacao controlada."),
        ("LGPD", "Lei Geral de Protecao de Dados.")
    ])

    return doc


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
